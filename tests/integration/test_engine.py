"""End-to-end runs of the engine on real files, including failure scenarios (phases 2-5)."""

from __future__ import annotations

import json
import os
import uuid
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from data_automation import engine, writers
from data_automation.engine import RunRejected, run_workflow

SALES_NORTH = (
    "region,order_id,status,amount\nNorth,SO-1,closed,120.50\nNorth,SO-2,open,80.00\nNorth,SO-3,closed,42.00\n"
)
SALES_SOUTH = "region,order_id,status,amount\nSouth,SO-4,closed,200.00\nSouth,SO-5,cancelled,15.00\n"
TARGETS = "region,target\nNorth,150.00\nSouth,250.00\n"

WORKFLOW = """\
workflow: {{id: monthly, schema_version: 1}}
parameters:
  period: {{type: month}}
inputs:
  - id: sales_files
    type: csv
    path: "in/{{period}}/sales_*.csv"
    required_columns: [region, order_id, status, amount]
    types: {{amount: decimal}}
  - id: targets
    type: csv
    path: ref/targets.csv
    types: {{target: decimal}}
steps:
  - {{id: consolidate, type: concatenate, input: sales_files, output: sales, add_source_column: source_file}}
  - id: closed
    type: filter
    input: sales
    output: closed_sales
    where: [{{column: status, op: eq, value: closed}}]
  - id: kpis
    type: aggregate
    input: closed_sales
    output: kpis
    group_by: [region]
    metrics:
      total_sales: {{function: sum, column: amount}}
      order_count: {{function: count}}
  - {{id: with_targets, type: join, input: [kpis, targets], output: report, keys: [region], relationship: one_to_one, how: left}}
outputs:
  - {{id: csv, input: report, type: csv, path: "out/sales_{{period}}.csv", expected_columns: [region, total_sales, order_count, target]}}
  - {{id: xlsx, input: report, type: xlsx, path: "out/sales_{{period}}.xlsx", sheet_name: KPIs}}
on_error: {on_error}
report:
  summary: {{datasets: [report], title: Monthly sales}}
"""

EXPECTED_CSV = "region,total_sales,order_count,target\nNorth,162.50,2,150.00\nSouth,200.00,1,250.00\n"


@pytest.fixture
def project(tmp_path):
    def _make(on_error="stop", workflow=None, north=SALES_NORTH):
        files = {
            "in/2026-09/sales_north.csv": north,
            "in/2026-09/sales_south.csv": SALES_SOUTH,
            "ref/targets.csv": TARGETS,
        }
        for relative, content in files.items():
            path = tmp_path / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(content, encoding="utf-8")
        wf = tmp_path / "workflow.yaml"
        wf.write_text(workflow or WORKFLOW.format(on_error=on_error), encoding="utf-8")
        return wf

    return _make


def _run(wf, **kwargs):
    return run_workflow(wf, {"period": "2026-09"}, **kwargs)


def _files(root: Path) -> dict[str, bytes]:
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in sorted(root.rglob("*")) if p.is_file()}


def test_full_run_success(project, tmp_path):
    result = _run(project())
    assert (result.status, result.exit_code) == ("success", 0)
    assert uuid.UUID(result.run_id).version == 4
    assert (tmp_path / "out/sales_2026-09.csv").read_text(encoding="utf-8") == EXPECTED_CSV
    sheet = load_workbook(tmp_path / "out/sales_2026-09.xlsx")["KPIs"]
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [
        ["region", "total_sales", "order_count", "target"],
        ["North", 162.5, 2, 150],
        ["South", 200, 1, 250],
    ]
    report = result.report
    assert [s["status"] for s in report["steps"]] == ["success"] * 4
    assert [s.get("rows_out") for s in report["steps"]] == [5, 3, 2, 2]
    assert [o["status"] for o in report["outputs"]] == ["published", "published"]
    assert all(s["unchanged"] for s in report["sources"]) and len(report["sources"]) == 3
    assert report["run_folder"] == f"runs/monthly/{result.run_id}"
    assert sorted(p.name for p in result.run_dir.iterdir()) == [
        "run.log",
        "run_report.json",
        "run_report.md",
        "summary.md",
    ]
    assert json.loads((result.run_dir / "run_report.json").read_text(encoding="utf-8"))["status"] == "success"
    summary = (result.run_dir / "summary.md").read_text(encoding="utf-8")
    assert "# Monthly sales" in summary and "| North | 162.50 | 2 | 150.00 |" in summary
    assert all(b"\r\n" not in p.read_bytes() for p in result.run_dir.iterdir())  # LF everywhere, as the outputs
    log = (result.run_dir / "run.log").read_text(encoding="utf-8")
    assert "Step 'kpis' (aggregate)" in log and "finished: success" in log
    assert not (tmp_path / "runs/monthly/.lock").exists()
    assert not list(tmp_path.rglob("*.tmp-*"))


def test_metadata_describes_the_published_file(project, tmp_path):
    result = _run(project())
    meta = json.loads((tmp_path / "out/sales_2026-09.csv.metadata.json").read_text(encoding="utf-8"))
    assert meta["run_id"] == result.run_id and meta["workflow_id"] == "monthly"
    assert meta["file"]["sha256"] == writers.sha256(tmp_path / "out/sales_2026-09.csv")
    assert meta["file"]["format"] == {
        "encoding": "utf-8",
        "bom": False,
        "delimiter": ",",
        "decimal_separator": ".",
        "line_terminator": "LF",
    }
    assert meta["columns"][1] == {"name": "total_sales", "type": "decimal", "nullable": False}
    assert [s["path"] for s in meta["sources"]] == [
        "in/2026-09/sales_north.csv",
        "in/2026-09/sales_south.csv",
        "ref/targets.csv",
    ]


def test_runs_are_repeatable_without_duplicates(project, tmp_path):
    wf = project()
    _run(wf)
    first = (tmp_path / "out/sales_2026-09.csv").read_bytes()
    second = _run(wf, overwrite=True)
    assert second.status == "success"
    assert (tmp_path / "out/sales_2026-09.csv").read_bytes() == first  # same data, same bytes, no appended rows
    assert len(list((tmp_path / "runs/monthly").iterdir())) == 2  # one run folder per run


def test_existing_outputs_are_never_overwritten_without_permission(project, tmp_path):
    wf = project()
    _run(wf)
    before = _files(tmp_path / "out")
    with pytest.raises(RunRejected) as info:
        _run(wf)
    assert len(info.value.problems) == 4  # two outputs + two metadata files, all listed at once
    assert _files(tmp_path / "out") == before
    assert len(list((tmp_path / "runs/monthly").iterdir())) == 1  # no run folder for a refused run


def test_stop_policy_publishes_nothing_and_keeps_previous_outputs(project, tmp_path):
    wf = project()
    _run(wf)
    before = _files(tmp_path / "out")
    (tmp_path / "in/2026-09/sales_north.csv").write_text(
        "region,order_id,status,amount\nNorth,SO-1,closed,abc\n", encoding="utf-8"
    )
    result = _run(wf, overwrite=True)
    assert (result.status, result.exit_code) == ("failed", 1)
    assert result.report["inputs"][0]["status"] == "failed"
    assert "row 2: 'abc'" in result.report["inputs"][0]["error"]
    assert [s["status"] for s in result.report["steps"]] == ["skipped"] * 4
    assert [o["status"] for o in result.report["outputs"]] == ["not_published", "not_published"]
    assert _files(tmp_path / "out") == before  # UC-04: previous results preserved


def test_continue_policy_gives_a_partial_run(project, tmp_path):
    workflow = WORKFLOW.format(on_error="continue")
    workflow = workflow.replace(
        "  - {id: with_targets",
        "  - {id: broken, type: cast, input: sales, output: never, columns: {order_id: integer}}\n  - {id: with_targets",
    )
    workflow = workflow.replace(
        "on_error: continue", "  - {id: never_out, input: never, type: csv, path: out/never.csv}\non_error: continue"
    )
    result = _run(project(workflow=workflow))
    assert (result.status, result.exit_code) == ("partial", 4)
    statuses = {s["id"]: s["status"] for s in result.report["steps"]}
    assert statuses == {
        "consolidate": "success",
        "closed": "success",
        "kpis": "success",
        "broken": "failed",
        "with_targets": "success",
    }
    assert "'order_id' are not valid integer" in result.report["steps"][3]["error"]
    outputs = {o["id"]: o["status"] for o in result.report["outputs"]}
    assert outputs == {"csv": "published", "xlsx": "published", "never_out": "skipped"}
    assert (tmp_path / "out/sales_2026-09.csv").read_text(encoding="utf-8") == EXPECTED_CSV
    assert not (tmp_path / "out/never.csv").exists()


def test_join_relationship_violation_fails_the_run(project, tmp_path):
    wf = project()
    (tmp_path / "ref/targets.csv").write_text(TARGETS + "North,999\n", encoding="utf-8")
    result = _run(wf)
    assert result.status == "failed"
    assert "relationship one_to_one violated" in result.report["steps"][3]["error"]
    assert not (tmp_path / "out").exists()


def test_validation_errors_reject_the_run_before_anything(project, tmp_path):
    wf = project(workflow=WORKFLOW.format(on_error="stop").replace("function: sum", "function: total"))
    with pytest.raises(RunRejected) as info:
        _run(wf)
    assert any("metrics.total_sales.function" in p for p in info.value.problems)
    assert not (tmp_path / "runs").exists() and not (tmp_path / "out").exists()


def test_lock_prevents_concurrent_runs(project, tmp_path):
    wf = project()
    lock = tmp_path / "runs/monthly/.lock"
    lock.parent.mkdir(parents=True)
    lock.write_text("run_id=other pid=1\n", encoding="utf-8")
    with pytest.raises(RunRejected, match="another run of this workflow holds the lock"):
        _run(wf)
    assert lock.read_text(encoding="utf-8") == "run_id=other pid=1\n"  # never removed by the refused run
    assert not (tmp_path / "out").exists()


def test_source_changed_during_the_run_blocks_publication(project, tmp_path, monkeypatch):
    wf = project(on_error="continue")
    real_prepare = engine._Run.prepare_outputs

    def prepare_then_touch(self):
        prepared = real_prepare(self)
        (tmp_path / "ref/targets.csv").write_text(TARGETS + "West,1\n", encoding="utf-8")
        return prepared

    monkeypatch.setattr(engine._Run, "prepare_outputs", prepare_then_touch)
    result = _run(wf)
    assert result.status == "failed"  # even with on_error: continue
    assert any("ref/targets.csv changed during the run" in e for e in result.report["errors"])
    assert {o["status"] for o in result.report["outputs"]} == {"not_published"}
    assert not (tmp_path / "out/sales_2026-09.csv").exists()
    assert not list(tmp_path.rglob("*.tmp-*"))


class _Raising:
    """Stands in for a step type whose run raises the given exception."""

    def __init__(self, exc):
        self.exc = exc

    def run(self, params, inputs, ctx):
        raise self.exc


def _replace_step(monkeypatch, name, exc):
    real_get = engine.get
    monkeypatch.setattr(engine, "get", lambda step_type: _Raising(exc) if step_type == name else real_get(step_type))


def test_interruption_publishes_nothing_and_cleans_up(project, tmp_path, monkeypatch):
    wf = project()
    _run(wf)
    before = _files(tmp_path / "out")
    _replace_step(monkeypatch, "aggregate", KeyboardInterrupt())
    result = _run(wf, overwrite=True)
    assert (result.status, result.exit_code) == ("interrupted", 130)
    assert [s["status"] for s in result.report["steps"]] == ["success", "success", "interrupted", "skipped"]
    assert {o["status"] for o in result.report["outputs"]} == {"not_published"}
    assert _files(tmp_path / "out") == before
    assert not (tmp_path / "runs/monthly/.lock").exists()
    assert json.loads((result.run_dir / "run_report.json").read_text(encoding="utf-8"))["status"] == "interrupted"
    assert not list(tmp_path.rglob("*.tmp-*"))


def test_disk_error_while_writing_fails_cleanly(project, tmp_path, monkeypatch):
    wf = project()

    def disk_full(dataset, path, sheet):
        path.write_bytes(b"partial")
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(engine, "write_xlsx", disk_full)
    result = _run(wf)
    assert result.status == "failed"
    assert "No space left on device" in result.report["outputs"][1]["error"]
    assert not (tmp_path / "out/sales_2026-09.csv").exists()  # stop: the verified CSV is not published either
    assert not list(tmp_path.rglob("*.tmp-*"))


def test_verification_failure_is_not_published(project, tmp_path, monkeypatch):
    wf = project(on_error="continue")

    def corrupt_then_verify(dataset, path):
        path.write_bytes(path.read_bytes().replace(b"162.50", b"162.5"))
        writers.verify_csv(dataset, path)

    monkeypatch.setattr(engine, "verify_csv", corrupt_then_verify)
    result = _run(wf)
    assert result.status == "partial"
    assert result.report["outputs"][0]["status"] == "failed" and "read back" in result.report["outputs"][0]["error"]
    assert not (tmp_path / "out/sales_2026-09.csv").exists()
    assert (tmp_path / "out/sales_2026-09.xlsx").exists()


@pytest.mark.skipif(os.name != "nt", reason="Windows file locking")
def test_output_open_in_another_program_is_reported(project, tmp_path):
    wf = project()
    _run(wf)
    with (tmp_path / "out/sales_2026-09.csv").open("rb"):
        result = _run(wf, overwrite=True)
    assert result.status == "failed"
    assert "could not be published" in " ".join(result.report["errors"])


def test_internal_step_error_is_recorded_not_hidden(project, monkeypatch):
    _replace_step(monkeypatch, "filter", ZeroDivisionError("bug"))
    result = _run(project())
    assert result.status == "failed"
    assert result.report["errors"][0] == "step 'closed' (filter): internal error ZeroDivisionError('bug')"


def test_step_failure_message_reaches_the_report(project):
    text = WORKFLOW.format(on_error="stop").replace(
        "  - {id: with_targets",
        "  - {id: guard, type: check, input: kpis, output: checked, min_rows: 5}\n  - {id: with_targets",
    )
    result = _run(project(workflow=text))
    assert result.report["steps"][3]["error"] == "check failed: 2 row(s), at least 5 expected"
    assert result.report["steps"][4]["status"] == "skipped"


def test_xlsx_input_and_empty_match_warning(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append(["sku", "qty", "price"])
    sheet.append(["AB-001", 2, 19.99])
    sheet.append(["AB-002", 1, 5])
    workbook.save(tmp_path / "orders.xlsx")
    (tmp_path / "w.yaml").write_text(
        "workflow: {id: orders, schema_version: 1}\n"
        "inputs:\n"
        "  - {id: orders, type: xlsx, path: orders.xlsx, sheet: Orders, types: {qty: integer, price: decimal}}\n"
        "  - {id: late, type: csv, path: 'late/*.csv', required_columns: [sku], on_empty_match: warn}\n"
        "steps:\n"
        "  - {id: total, type: aggregate, input: orders, output: totals, metrics: {revenue: {function: sum, column: price}, units: {function: sum, column: qty}}}\n"
        "  - {id: late_count, type: aggregate, input: late, output: late_n, metrics: {n: {function: count}}}\n"
        "outputs:\n"
        "  - {id: out, input: totals, type: xlsx, path: totals.xlsx}\n"
        "  - {id: late_out, input: late_n, type: csv, path: late.csv}\n",
        encoding="utf-8",
    )
    result = run_workflow(tmp_path / "w.yaml")
    assert result.status == "success"
    rows = [[c.value for c in r] for r in load_workbook(tmp_path / "totals.xlsx")["data"].iter_rows()]
    assert rows == [["revenue", "units"], [24.99, 3]]
    assert (tmp_path / "late.csv").read_text(encoding="utf-8") == "n\n0\n"
    assert any("no file matches late/*.csv" in w for w in result.report["warnings"])


def test_period_parameter_selects_only_that_period_files(project, tmp_path):
    """AC-18: the run parameter chooses the files; other periods are never read."""
    wf = project()
    august = tmp_path / "in/2026-08"
    august.mkdir(parents=True)
    (august / "sales_north.csv").write_text(
        "region,order_id,status,amount\nNorth,SO-9,closed,1000.00\n", encoding="utf-8"
    )
    result = run_workflow(wf, {"period": "2026-08"})
    assert result.status == "success"
    assert [s["path"] for s in result.report["sources"]] == ["in/2026-08/sales_north.csv", "ref/targets.csv"]
    assert (tmp_path / "out/sales_2026-08.csv").read_text(encoding="utf-8") == (
        "region,total_sales,order_count,target\nNorth,1000.00,1,150.00\n"
    )
    assert not (tmp_path / "out/sales_2026-09.csv").exists()
    with pytest.raises(RunRejected, match="must be a month"):
        run_workflow(wf, {"period": "2026-8"})
