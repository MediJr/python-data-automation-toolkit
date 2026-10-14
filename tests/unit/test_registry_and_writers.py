"""Step registry extensibility (AC-12) and the export contract (ADR-0007)."""

from __future__ import annotations

import datetime as dt
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from data_automation import registry, writers
from data_automation.dataset import Dataset
from data_automation.engine import run_workflow
from data_automation.registry import StepParams, StepResult, StepType
from data_automation.writers import ExportError, check_columns, verify_csv, verify_xlsx, write_csv, write_xlsx

# --------------------------------------------------------------------------- registry (AC-12)


class UpperParams(StepParams):
    column: str


def _run_upper(params, inputs, ctx):
    frame = inputs[0].frame().copy()
    frame[params.column] = frame[params.column].str.upper()
    return StepResult(Dataset.from_frame(ctx.step_id, frame, inputs[0].types), ["uppercased"], {})


@pytest.fixture
def upper_step():
    step = registry.register(StepType("upper", "Upper-case one text column (test-only).", UpperParams, _run_upper))
    yield step
    registry.unregister("upper")


def test_new_step_type_runs_through_the_unchanged_engine(upper_step, tmp_path):
    (tmp_path / "in.csv").write_text("city\nparis\nlyon\n", encoding="utf-8")
    (tmp_path / "w.yaml").write_text(
        "workflow: {id: custom, schema_version: 1}\n"
        "inputs: [{id: cities, type: csv, path: in.csv}]\n"
        "steps: [{id: shout, type: upper, input: cities, output: loud, column: city}]\n"
        "outputs: [{id: out, input: loud, type: csv, path: out.csv}]\n",
        encoding="utf-8",
    )
    result = run_workflow(tmp_path / "w.yaml")
    assert result.status == "success"
    assert (tmp_path / "out.csv").read_text(encoding="utf-8") == "city\nPARIS\nLYON\n"
    assert result.report["steps"][0]["warnings"] == ["uppercased"]


def test_registry_refuses_duplicates_and_validates_new_types(upper_step):
    with pytest.raises(ValueError, match="already registered"):
        registry.register(upper_step)
    _, problems = registry.validate_step("upper", 1, {"column": 3, "extra": 1})
    assert "column: Input should be a valid string (got 3)" in problems
    assert any(p.startswith("extra: unknown key") for p in problems)
    assert "upper" in registry.available()


# --------------------------------------------------------------------------- writers

DATA = Dataset.from_columns(
    "kpis",
    {
        "code": ["007", "=SUM(A1)", None],
        "amount": [Decimal("1234.50"), Decimal("-0.10"), None],
        "count": [3, None, 0],
        "day": [dt.date(2026, 9, 30), None, dt.date(2026, 1, 1)],
        "at": [dt.datetime(2026, 9, 30, 14, 5), None, None],
        "flag": [True, False, None],
        "note": ['say "hi", ok', "multi\nline", "Zürich"],
    },
    {
        "code": "string",
        "amount": "decimal",
        "count": "integer",
        "day": "date",
        "at": "datetime",
        "flag": "boolean",
        "note": "string",
    },
)

EXPECTED_CSV = (
    "code,amount,count,day,at,flag,note\n"
    '007,1234.50,3,2026-09-30,2026-09-30T14:05:00,true,"say ""hi"", ok"\n'
    '=SUM(A1),-0.10,,,,false,"multi\nline"\n'
    ",,0,2026-01-01,,,Zürich\n"
)


def test_csv_contract_exact_bytes(tmp_path):
    path = tmp_path / "out.csv"
    write_csv(DATA, path)
    raw = path.read_bytes()
    assert raw == EXPECTED_CSV.encode("utf-8")  # UTF-8 without BOM, LF, ISO dates, '.' decimals, empty = missing
    assert not raw.startswith(b"\xef\xbb\xbf") and b"\r\n" not in raw
    verify_csv(DATA, path)


def test_csv_verification_detects_any_difference(tmp_path):
    path = tmp_path / "out.csv"
    write_csv(DATA, path)
    path.write_bytes(path.read_bytes().replace(b"1234.50", b"1234.5"))
    with pytest.raises(ExportError, match="data row 1: read back"):
        verify_csv(DATA, path)
    path.write_bytes(b"\xef\xbb\xbf" + EXPECTED_CSV.encode("utf-8"))
    with pytest.raises(ExportError, match="BOM"):
        verify_csv(DATA, path)


def test_xlsx_typed_cells_and_verification(tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(DATA, path, "data")
    verify_xlsx(DATA, path, "data")
    sheet = load_workbook(path)["data"]
    assert sheet["A2"].value == "007" and sheet["A2"].data_type == "s"
    assert sheet["A3"].value == "=SUM(A1)" and sheet["A3"].data_type == "s"  # text, never a formula
    assert sheet["B2"].value == 1234.5 and sheet["B2"].data_type == "n"
    assert sheet["D2"].is_date and sheet["D2"].number_format == "yyyy-mm-dd"
    assert sheet["F2"].value is True
    assert sheet["B4"].value is None


def test_xlsx_verification_detects_a_changed_cell(tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(DATA, path, "data")
    workbook = load_workbook(path)
    workbook["data"]["C2"] = 4
    workbook.save(path)
    with pytest.raises(ExportError, match="data row 1, column 'count': read back 4 instead of 3"):
        verify_xlsx(DATA, path, "data")


def test_xlsx_refuses_control_characters_and_time_zones(tmp_path):
    bad = Dataset.from_columns("b", {"t": ["a\x01b"]}, {"t": "string"})
    with pytest.raises(ExportError, match="control character"):
        write_xlsx(bad, tmp_path / "b.xlsx", "data")
    aware = Dataset.from_columns("a", {"t": [dt.datetime(2026, 1, 1, tzinfo=dt.UTC)]}, {"t": "datetime"})
    with pytest.raises(ExportError, match="time zones"):
        write_xlsx(aware, tmp_path / "a.xlsx", "data")


def test_xlsx_refuses_numbers_excel_would_round(tmp_path):
    exact = Dataset.from_columns("e", {"d": [Decimal("0.1"), Decimal("123456789012.34")]}, {"d": "decimal"})
    write_xlsx(exact, tmp_path / "ok.xlsx", "data")
    verify_xlsx(exact, tmp_path / "ok.xlsx", "data")
    long = Dataset.from_columns("l", {"d": [Decimal("1.00000000000000000001")]}, {"d": "decimal"})
    with pytest.raises(ExportError, match="cannot be stored exactly as an Excel number"):
        write_xlsx(long, tmp_path / "l.xlsx", "data")
    big = Dataset.from_columns("b", {"n": [2**53 + 1]}, {"n": "integer"})
    with pytest.raises(ExportError, match="too large"):
        write_xlsx(big, tmp_path / "b.xlsx", "data")
    assert not (tmp_path / "l.xlsx").exists() and not (tmp_path / "b.xlsx").exists()


def test_expected_columns_must_match_exactly():
    check_columns(DATA, None)
    check_columns(DATA, list(DATA.columns))
    with pytest.raises(ExportError, match=r"missing \['total'\]"):
        check_columns(DATA, ["code", "total"])
    with pytest.raises(ExportError, match="order is"):
        check_columns(DATA, list(reversed(DATA.columns)))


def test_metadata_helpers():
    assert writers.decimal_places(DATA) == {"amount": 2}
    schema = writers.column_schema(DATA)
    assert schema[0] == {"name": "code", "type": "string", "nullable": True}
    assert {c["name"]: c["nullable"] for c in schema}["note"] is False
