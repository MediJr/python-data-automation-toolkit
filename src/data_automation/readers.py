"""Reading source files into datasets (phase 3).

Rules, identical for CSV and XLSX:

* every cell is first read as text; columns are converted only to the types declared in
  the input (``types``); other columns stay ``string``. Nothing is guessed;
* an empty cell is a missing value (never 0 or an empty default);
* a value that does not match its declared type fails the input, with the file, the
  row and the value;
* each file becomes one part of the dataset, so that files are combined explicitly by a
  ``concatenate`` step; ``required_columns`` and declared columns must exist in every file;
* structural problems (rows with the wrong number of fields, empty or duplicate header
  names, undecodable bytes, Excel formulas without a saved value) fail the input.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import pandas as pd
from defusedxml import ElementTree  # refuses entity declarations: .xlsx files are untrusted input
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .config import InputSpec
from .dataset import Dataset, Part
from .paths import relative_posix
from .types import parse_text, to_series

MAX_LISTED = 5


class InputReadError(Exception):
    """A source file cannot be used; the message names the file and the problem."""


@dataclass
class ReadResult:
    dataset: Dataset
    warnings: list[str] = field(default_factory=list)
    files: list[dict] = field(default_factory=list)


@dataclass
class _Table:
    header: list[str]
    rows: list[tuple[int, list[str]]]  # (row/line number in the file, cell texts)
    skipped_blank_rows: int = 0


# --------------------------------------------------------------------------- public API


def read_input(spec: InputSpec, files: list[Path], root: Path) -> ReadResult:
    """Read the files of an input into a dataset with one part per file.

    With no file at all (only possible with ``on_empty_match: warn``), the result is an
    empty dataset whose columns are the declared ones.

    Raises:
        InputReadError: Any file cannot be read or does not match the declaration.
    """
    warnings: list[str] = []
    parts: list[Part] = []
    types: dict[str, str] = {}
    file_info: list[dict] = []
    for path in files:
        name = relative_posix(root, path)
        if _is_empty_file(path):
            if spec.on_empty_file == "skip":
                warnings.append(f"{name} is empty and was skipped (on_empty_file: skip)")
                file_info.append({"path": name, "rows": 0, "skipped": True})
                continue
            raise InputReadError(f"{name} is empty (set on_empty_file: skip to ignore empty files)")
        table = _read_csv(path, spec, name) if spec.type == "csv" else _read_xlsx(path, spec, name)
        if table.skipped_blank_rows:
            warnings.append(f"{name}: {table.skipped_blank_rows} completely empty row(s) ignored")
        frame, file_types = _typed_frame(table, spec, name)
        for column, kind in file_types.items():
            types.setdefault(column, kind)
        parts.append(Part(frame, name))
        file_info.append({"path": name, "rows": len(frame), "columns": list(frame.columns)})
    if files and not parts:
        raise InputReadError(f"all {len(files)} matching file(s) are empty")
    if not files:
        columns = list(dict.fromkeys([*spec.required_columns, *spec.types]))
        types = {c: spec.types[c].type if c in spec.types else "string" for c in columns}
        empty = pd.DataFrame({c: to_series(types[c], []) for c in columns})
        return ReadResult(Dataset(spec.id, columns, types, [Part(empty, spec.id)]), warnings, file_info)
    return ReadResult(Dataset(spec.id, list(parts[0].frame.columns), types, parts), warnings, file_info)


# --------------------------------------------------------------------------- typing


def _typed_frame(table: _Table, spec: InputSpec, name: str) -> tuple[pd.DataFrame, dict[str, str]]:
    missing = [c for c in [*spec.required_columns, *spec.types] if c not in table.header]
    if missing:
        raise InputReadError(
            f"{name}: missing column(s) {', '.join(repr(c) for c in dict.fromkeys(missing))} "
            f"(found: {', '.join(repr(c) for c in table.header)})"
        )
    columns: dict[str, pd.Series] = {}
    types: dict[str, str] = {}
    problems = []
    for position, column in enumerate(table.header):
        declared = spec.types.get(column)
        kind = declared.type if declared else "string"
        values, bad = [], []
        for number, cells in table.rows:
            text = cells[position]
            if text == "" or (kind != "string" and text.strip() == ""):
                values.append(None)
                continue
            try:
                values.append(
                    parse_text(
                        kind,
                        text,
                        date_format=declared.format if declared else None,
                        decimal_separator=spec.decimal_separator,
                    )
                )
            except ValueError as exc:
                values.append(None)
                bad.append((number, text, str(exc)))
        if bad:
            listed = "; ".join(f"row {n}: {t!r}" for n, t, _ in bad[:MAX_LISTED])
            more = f" and {len(bad) - MAX_LISTED} more" if len(bad) > MAX_LISTED else ""
            problems.append(
                f"column {column!r}: {len(bad)} value(s) are not valid {kind} ({bad[0][2]}): {listed}{more}"
            )
        columns[column] = to_series(kind, values)
        types[column] = kind
    if problems:
        raise InputReadError(f"{name}: " + " | ".join(problems))
    return pd.DataFrame(columns), types


def _check_header(header: list[str], name: str) -> list[str]:
    blanks = [str(i + 1) for i, h in enumerate(header) if h.strip() == ""]
    if blanks:
        raise InputReadError(f"{name}: column(s) {', '.join(blanks)} have no name in the header")
    duplicates = sorted({h for h in header if header.count(h) > 1})
    if duplicates:
        raise InputReadError(f"{name}: duplicate column name(s) {duplicates}")
    return header


def _is_empty_file(path: Path) -> bool:
    if path.stat().st_size == 0:
        return True
    if path.suffix.lower() == ".csv" and path.stat().st_size < 4096:
        return path.read_bytes().strip(b" \t\r\n\xef\xbb\xbf") == b""
    return False


# --------------------------------------------------------------------------- CSV


def _read_csv(path: Path, spec: InputSpec, name: str) -> _Table:
    raw = path.read_bytes()
    encoding = "utf-8-sig" if spec.encoding.lower().replace("_", "-") in ("utf-8", "utf8") else spec.encoding
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError as exc:
        raise InputReadError(
            f"{name} is not valid {spec.encoding} (byte {exc.start}); set the input's encoding"
        ) from None
    if "\x00" in text:
        raise InputReadError(f"{name} contains NUL bytes; it is probably not a {spec.encoding} CSV file")
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=spec.delimiter, strict=True)
    records: list[tuple[int, list[str]]] = []
    previous = 0
    try:
        for record in reader:
            start, previous = previous + 1, reader.line_num
            if record:
                records.append((start, record))
    except csv.Error as exc:
        raise InputReadError(f"{name}: invalid CSV near line {reader.line_num}: {exc}") from None
    if not records:
        raise InputReadError(f"{name} has no header")
    header = _check_header(records[0][1], name)
    bad = [(n, len(r)) for n, r in records[1:] if len(r) != len(header)]
    if bad:
        listed = ", ".join(f"line {n} ({c} fields)" for n, c in bad[:MAX_LISTED])
        raise InputReadError(
            f"{name}: the header has {len(header)} fields but {len(bad)} row(s) do not: {listed}; check the delimiter"
        )
    rows = [(n, r) for n, r in records[1:] if any(cell != "" for cell in r)]
    return _Table(header, rows, len(records) - 1 - len(rows))


# --------------------------------------------------------------------------- XLSX

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _excel_text(value: object, decimal_separator: str = ".") -> str:
    """Text of a typed cell, in the form the type parsers accept.

    Numbers are written without exponent (``1e-05`` becomes ``0.00001``) and with the
    input's decimal separator, so a numeric cell reads the same as the text a user would
    have typed in that file.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e16:
            return str(int(value))
        return format(Decimal(repr(value)), "f").replace(".", decimal_separator)
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    return str(value)


def _uncalculated_formulas(path: Path, sheet: str) -> int | None:
    """Formula cells of ``sheet`` that have no saved value (``None`` if not inspectable)."""
    try:
        with zipfile.ZipFile(path) as archive:
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {rel.get("Id"): rel.get("Target", "") for rel in rels}
            rel_id = next(s.get(f"{_REL}id") for s in workbook.iter(f"{_NS}sheet") if s.get("name") == sheet)
            target = targets[rel_id]
            member = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            missing = 0
            with archive.open(member) as handle:
                for _, element in ElementTree.iterparse(handle):
                    if element.tag == f"{_NS}c":
                        if element.find(f"{_NS}f") is not None:
                            value = element.find(f"{_NS}v")
                            missing += value is None or not value.text
                        element.clear()
                    elif element.tag == f"{_NS}row":
                        element.clear()
            return missing
    except KeyError, StopIteration, zipfile.BadZipFile, ElementTree.ParseError, DefusedXmlException, OSError:
        return None


def _read_xlsx(path: Path, spec: InputSpec, name: str) -> _Table:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError, DefusedXmlException) as exc:
        raise InputReadError(
            f"{name} cannot be opened as an .xlsx workbook "
            f"(corrupt, password-protected, unsafe XML or not .xlsx: {type(exc).__name__})"
        ) from None
    try:
        names = list(workbook.sheetnames)
        sheet = spec.sheet or names[0]
        if sheet not in names:
            raise InputReadError(f"{name}: sheet {sheet!r} not found (available: {', '.join(names)})")
        worksheet = workbook[sheet]
        if not hasattr(worksheet, "iter_rows"):
            raise InputReadError(f"{name}: sheet {sheet!r} is a chart sheet")
        raw_rows = [
            (i, [_excel_text(v, spec.decimal_separator) for v in row])
            for i, row in enumerate(worksheet.iter_rows(values_only=True), 1)
        ]
    finally:
        workbook.close()
    uncalculated = _uncalculated_formulas(path, sheet)
    if uncalculated:
        raise InputReadError(
            f"{name}: {uncalculated} formula cell(s) on sheet {sheet!r} have no saved value (the file was never "
            "recalculated by Excel); open and save it in Excel first"
        )
    non_empty = [(i, r) for i, r in raw_rows if any(r)]
    if not non_empty:
        raise InputReadError(f"{name}: sheet {sheet!r} has no header")
    header_row, header = non_empty[0]
    while header and header[-1] == "":
        header = header[:-1]
    _check_header(header, name)
    width = len(header)
    rows = []
    for number, cells in raw_rows:
        if number <= header_row:
            continue
        if any(cells[width:]):
            raise InputReadError(f"{name}: row {number} has values outside the {width} header column(s)")
        cells = (cells + [""] * width)[:width]
        if any(cells):
            rows.append((number, cells))
    skipped = sum(1 for n, c in raw_rows if n > header_row and not any(c[:width]))
    trailing = 0  # trailing empty rows of a formatted sheet are not worth a warning
    for _n, cells in reversed([(n, c) for n, c in raw_rows if n > header_row]):
        if any(cells[:width]):
            break
        trailing += 1
    return _Table(header, rows, skipped - trailing)
