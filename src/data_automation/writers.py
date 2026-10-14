"""Export contract (ADR-0007): CSV and XLSX files, verified by reading them back, plus metadata.

CSV: UTF-8 **without BOM**, ``,`` delimiter, standard quoting, ``LF`` line endings, one
header row, ISO 8601 dates, ``.`` as decimal separator without grouping, empty cells for
missing values.

XLSX: one sheet, a header row, numbers as numeric cells (decimals keep their exact text in
the file), dates as date cells shown ``yyyy-mm-dd``, text as text (a value starting with
``=`` never becomes a formula).

Each file is written to a temporary name next to its destination, read back completely
and compared with the dataset; only then may the engine publish it.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .dataset import Dataset
from .types import format_text, python_values

CSV_FORMAT = {"encoding": "utf-8", "bom": False, "delimiter": ",", "decimal_separator": ".", "line_terminator": "LF"}
DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"


class ExportError(Exception):
    """The export could not be written or did not read back identically."""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def check_columns(dataset: Dataset, expected: list[str] | None) -> None:
    """``expected_columns`` must match the dataset's columns exactly, in order."""
    if expected is None or dataset.columns == expected:
        return
    missing = [c for c in expected if c not in dataset.columns]
    extra = [c for c in dataset.columns if c not in expected]
    detail = []
    if missing:
        detail.append(f"missing {missing}")
    if extra:
        detail.append(f"unexpected {extra}")
    if not detail:
        detail.append(f"order is {dataset.columns}")
    raise ExportError(f"columns do not match expected_columns {expected}: {'; '.join(detail)}")


def _rows(dataset: Dataset) -> list[list[Any]]:
    frame = dataset.frame()
    columns = [python_values(frame[c]) for c in dataset.columns]
    return [list(row) for row in zip(*columns, strict=True)] if columns else []


# --------------------------------------------------------------------------- CSV


def _csv_text(dataset: Dataset) -> str:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=",", lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(dataset.columns)
    kinds = [dataset.types[c] for c in dataset.columns]
    for row in _rows(dataset):
        writer.writerow([format_text(kind, value) for kind, value in zip(kinds, row, strict=True)])
    return buffer.getvalue()


def write_csv(dataset: Dataset, path: Path) -> None:
    path.write_bytes(_csv_text(dataset).encode("utf-8"))


def verify_csv(dataset: Dataset, path: Path) -> None:
    raw = path.read_bytes()
    if raw.startswith(b"\xef\xbb\xbf"):
        raise ExportError("the CSV starts with a BOM")
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ExportError(f"the CSV is not valid UTF-8 (byte {exc.start})") from None
    expected = list(csv.reader(io.StringIO(_csv_text(dataset), newline="")))
    actual = list(csv.reader(io.StringIO(text, newline="")))
    if actual[:1] != expected[:1]:
        raise ExportError(f"header read back {actual[:1]} differs from {expected[:1]}")
    if len(actual) != len(expected):
        raise ExportError(f"read back {len(actual) - 1} data rows, expected {len(expected) - 1}")
    for number, (got, want) in enumerate(zip(actual, expected, strict=True)):
        if got != want:
            raise ExportError(f"data row {number}: read back {got} instead of {want}")


# --------------------------------------------------------------------------- XLSX


def _check_xlsx_values(dataset: Dataset, rows: list[list[Any]]) -> None:
    """Refuse values .xlsx cannot store, before any file is created (no half-written workbook)."""
    for name in dataset.columns:
        if ILLEGAL_CHARACTERS_RE.search(name):
            raise ExportError(f"column name {name!r} contains a control character that .xlsx cannot store")
    for number, row in enumerate(rows, start=1):
        for column, value in zip(dataset.columns, row, strict=True):
            kind = dataset.types[column]
            if kind == "string" and value is not None and ILLEGAL_CHARACTERS_RE.search(value):
                raise ExportError(
                    f"data row {number}, column {column!r} contains a control character that .xlsx cannot store; "
                    "export as CSV"
                )
            if kind == "datetime" and value is not None and value.tzinfo is not None:
                raise ExportError(f"data row {number}, column {column!r}: xlsx cannot store time zones; export as CSV")
            # Excel numbers are binary doubles: refuse values they would round instead of changing them silently
            if kind == "decimal" and value is not None and Decimal(repr(float(value))) != value:
                raise ExportError(
                    f"data row {number}, column {column!r}: {value} cannot be stored exactly as an Excel number "
                    "(about 15 significant digits); export as CSV or round it"
                )
            if kind == "integer" and value is not None and abs(value) > 2**53:
                raise ExportError(
                    f"data row {number}, column {column!r}: {value} is too large to be stored exactly as an Excel "
                    "number; export as CSV"
                )


def write_xlsx(dataset: Dataset, path: Path, sheet_name: str) -> None:
    rows = _rows(dataset)
    _check_xlsx_values(dataset, rows)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(sheet_name)
    kinds = [dataset.types[c] for c in dataset.columns]
    header = []
    for name in dataset.columns:
        cell = WriteOnlyCell(sheet, value=name)
        cell.data_type = "s"
        header.append(cell)
    sheet.append(header)
    for row in rows:
        cells = []
        for kind, value in zip(kinds, row, strict=True):
            cell = WriteOnlyCell(sheet, value=value)
            if kind == "string" and value is not None:
                cell.data_type = "s"  # "=SUM(...)" stays text, never a formula
            elif kind == "date" and value is not None:
                cell.number_format = DATE_FORMAT
            elif kind == "datetime" and value is not None:
                cell.number_format = DATETIME_FORMAT
            cells.append(cell)
        sheet.append(cells)
    workbook.save(path)


def _same(kind: str, expected: Any, got: Any) -> bool:
    if expected is None:
        return got is None or got == ""
    if got is None:
        return False
    if kind == "string":
        return got == expected
    if kind == "boolean":
        return isinstance(got, bool) and got == expected
    if kind == "integer":
        return not isinstance(got, bool) and isinstance(got, int | float) and got == expected
    if kind == "decimal":
        return isinstance(got, int | float) and float(got) == float(expected)
    if kind == "date":
        got_date = got.date() if isinstance(got, dt.datetime) else got
        return got_date == expected
    if kind == "datetime":
        return isinstance(got, dt.datetime) and got.replace(microsecond=0) == expected.replace(microsecond=0)
    return False


def verify_xlsx(dataset: Dataset, path: Path, sheet_name: str) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any failure to reopen means the export is not verified
        raise ExportError(f"the written workbook cannot be reopened: {exc!r}") from None
    try:
        if sheet_name not in workbook.sheetnames:
            raise ExportError(f"sheet {sheet_name!r} not found when reading back")
        rows = [list(r) for r in workbook[sheet_name].iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not rows or [str(v) for v in rows[0]] != dataset.columns:
        raise ExportError(f"header read back {rows[:1]} differs from {dataset.columns}")
    expected = _rows(dataset)
    data = rows[1:]
    if len(data) != len(expected):
        raise ExportError(f"read back {len(data)} data rows, expected {len(expected)}")
    kinds = [dataset.types[c] for c in dataset.columns]
    for number, (got, want) in enumerate(zip(data, expected, strict=True), start=1):
        got = (got + [None] * len(kinds))[: len(kinds)]
        for column, kind, g, w in zip(dataset.columns, kinds, got, want, strict=True):
            if not _same(kind, w, g):
                raise ExportError(f"data row {number}, column {column!r}: read back {g!r} instead of {w!r}")


# --------------------------------------------------------------------------- metadata


def column_schema(dataset: Dataset) -> list[dict[str, Any]]:
    frame = dataset.frame()
    return [{"name": c, "type": dataset.types[c], "nullable": bool(frame[c].isna().any())} for c in dataset.columns]


def decimal_places(dataset: Dataset) -> dict[str, int]:
    """Largest number of decimal places per decimal column (useful for Power BI formats)."""
    frame = dataset.frame()
    places = {}
    for column, kind in dataset.types.items():
        if kind == "decimal":
            values = [v for v in python_values(frame[column]) if isinstance(v, Decimal)]
            places[column] = max((max(0, -v.as_tuple().exponent) for v in values), default=0)
    return places
